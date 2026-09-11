"""
Self-check for the Monitoring search / Excel export parity.

Guards against: slicing before the date filter (a QR found by search missing
from the export), export columns drifting behind the search payload, and
timestamps the export cannot split.

Run:  ./venv/bin/python test_monitoring_export.py
Read-only against the database.
"""
import os
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from datetime import datetime  # noqa: E402
from io import BytesIO  # noqa: E402

from django.test import RequestFactory  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from tracebility import models  # noqa: E402
from tracebility.monitoring_views import EXPORT_COLUMNS, monitoring_export_excel, search_monitoring_data  # noqa: E402
from tracebility.views import parse_timestamp_to_datetime  # noqa: E402

WIDE = {'time_filter': '100000hour'}


def _hours_back_to_newest_row():
    """Hours from now back to the newest prep row, +1 so it is inside the window."""
    newest = parse_timestamp_to_datetime(models.CiPreprocessing.objects.first().timestamp)
    return int((datetime.now() - newest).total_seconds() // 3600) + 1


def test_search_returns_every_row_not_a_1000_per_machine_sample():
    results = search_monitoring_data(WIDE)
    counts = {}
    for record in results:
        counts[record['display_name']] = counts.get(record['display_name'], 0) + 1
    capped = [m for m, n in counts.items() if n == 1000]
    assert not capped, f'still capped at exactly 1000: {capped}'
    total_ci = models.CiPreprocessing.objects.count()
    assert counts.get('Camera Inspection') == total_ci, f"CI: got {counts.get('Camera Inspection')} of {total_ci}"


def test_a_record_older_than_1000_rows_is_still_returned():
    old = models.CiPreprocessing.objects.all()[5000]
    by_qr = search_monitoring_data({'qr_code': old.qr_data, **WIDE})
    assert by_qr, f'QR search found nothing for {old.qr_data}'
    assert old.qr_data in {r['qr_code'] for r in search_monitoring_data(WIDE)}, 'missing from the broad result'


def test_every_field_the_search_returns_has_an_export_column():
    results = search_monitoring_data(WIDE)
    assert results, 'need rows or this check passes vacuously'
    fields = {k for r in results for k in r}
    exported = {
        'prep_id', 'post_id', 'display_name', 'machine_type', 'qr_code', 'timestamp', 'post_timestamp',
        'status', 'model_name', 'previous_machine_status', 'batch_id', 'batch_size', 'slot_no',
        'machine_name',  # internal id, covered by the Machine column
    }
    missing = fields - exported
    assert not missing, f'search returns fields with no export column: {sorted(missing)}'
    assert len(EXPORT_COLUMNS) == 15


def test_timestamps_are_all_parseable():
    bad = []
    for record in search_monitoring_data(WIDE):
        try:
            datetime.fromisoformat(record['timestamp'])
        except ValueError:
            bad.append(record)
    assert not bad, f'{len(bad)} records have a timestamp the export cannot split'


def test_export_sheet_has_exactly_the_search_rows():
    # A bounded window, but sized from the newest row so it always covers data
    # no matter how long ago the database was last written to.
    filters = {'time_filter': f'{_hours_back_to_newest_row()}hour'}
    results = search_monitoring_data(filters)
    assert results, 'window should contain data'
    request = RequestFactory().get('/monitoring/export-excel/', filters)
    ws = load_workbook(BytesIO(monitoring_export_excel(request).content)).active
    assert ws.max_row - 1 == len(results), f'sheet {ws.max_row - 1} rows vs search {len(results)}'
    assert [c.value for c in ws[1]] == [h for h, _ in EXPORT_COLUMNS]
    qrs_in_sheet = {ws.cell(row=i, column=5).value for i in range(2, ws.max_row + 1)}
    assert results[0]['qr_code'] in qrs_in_sheet


def test_model_and_status_filters_narrow_correctly():
    results = search_monitoring_data({'model_name': '68603833AA', 'status': 'NG', **WIDE})
    assert results
    assert all(r['model_name'] == '68603833AA' and r['status'] == 'NG' for r in results)


if __name__ == '__main__':
    failures = []
    for name, fn in sorted(globals().items()):
        if not name.startswith('test_'):
            continue
        try:
            fn()
            print(f'ok  {name}')
        except AssertionError as exc:
            failures.append(name)
            print(f'FAIL {name}: {exc}')
    print('\n' + ('All export checks passed.' if not failures else f'{len(failures)} FAILED'))
    raise SystemExit(1 if failures else 0)
